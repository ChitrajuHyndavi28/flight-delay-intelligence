"""Builds PROJECT_REPORT.docx from the university template 'SAMPLE FORMAT.docx' (cover, contents, abstract, lists, chapters 1-7).
The template's cover page, page borders, footers and table-of-contents table are kept; the body is replaced with this project's content."""
import copy, json, math, re, sys, textwrap
from pathlib import Path
from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent; SAMPLE = Path("/Users/hyndavi/Downloads/SAMPLE FORMAT.docx"); R = ROOT / "results"; SC = ROOT / "screenshots"; DATA = ROOT / "dashboard" / "data"
J = lambda n: json.load(open(DATA / f"{n}.json"))
doc = Document(SAMPLE); P = list(doc.paragraphs); body = doc.element.body
TITLE = "FLIGHT DELAY AND AIRPORT OPERATIONS INTELLIGENCE"

# ======================================================================= helpers ===============================================================
BM = {}; _bm_id = [200]; BM_TEXT = {}; est = {"page": 6, "line": 0.0}; LPP, CPL = 31.0, 86
FIGS, TABS = [], []          # (number, caption, bookmark)
def lines_of(text, cpl=CPL): return max(1, math.ceil(len(text) / cpl))
def use(lines, atomic=False):
    if atomic and est["line"] + lines > LPP: est["page"] += 1; est["line"] = 0.0
    est["line"] += lines
    while est["line"] > LPP: est["page"] += 1; est["line"] -= LPP
def new_page(): est["page"] += 1; est["line"] = 0.0
def bookmark(par, name):
    _bm_id[0] += 1; s = OxmlElement("w:bookmarkStart"); s.set(qn("w:id"), str(_bm_id[0])); s.set(qn("w:name"), name); e = OxmlElement("w:bookmarkEnd"); e.set(qn("w:id"), str(_bm_id[0]))
    ppr = par._p.find(qn("w:pPr")); (ppr.addnext(s) if ppr is not None else par._p.insert(0, s)); par._p.append(e); BM[name] = est["page"]
def runs(par, text, size, bold=False, italic=False, font=None, color=None):
    for i, chunk in enumerate(re.split(r"\*\*", text)):
        if not chunk: continue
        r = par.add_run(chunk); r.font.size = Pt(size); r.bold = bold or (i % 2 == 1); r.italic = italic
        if font: r.font.name = font; r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
        if color: r.font.color.rgb = RGBColor.from_string(color)
def para(text="", size=12, bold=False, italic=False, align="both", first=0.5, before=0, after=0, spacing=1.5, left=None, hanging=None, keep=False, brk=False, font=None):
    p = doc.add_paragraph(); pf = p.paragraph_format; pf.line_spacing = spacing; pf.space_before = Pt(before); pf.space_after = Pt(after)
    if first and not hanging: pf.first_line_indent = Inches(first)
    if left is not None: pf.left_indent = Inches(left)
    if hanging: pf.first_line_indent = Inches(-hanging)
    p.alignment = {"both": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT}[align]; pf.keep_with_next = keep; pf.page_break_before = brk
    runs(p, text, size, bold, italic, font); return p
def text(t, **k): p = para(t, **k); use(lines_of(t) * (k.get("spacing", 1.5) / 1.5) + 0.15); return p
def bullet(t, size=12): p = para("•\t" + t, size=size, first=0, left=0.55, hanging=0.25, after=2); p.paragraph_format.tab_stops.add_tab_stop(Inches(0.55)); use(lines_of(t, 80) + 0.1); return p
def chapter(n, title, bm):
    p = para(f"CHAPTER {n}", size=14, bold=True, align="center", first=0, brk=True); new_page(); use(1.2); p2 = para(title, size=14, bold=True, align="center", first=0, after=6); use(1.4); bookmark(p2, bm); BM_TEXT[bm] = f"CHAPTER {n}"; return p2
def h1(num, title, bm=None): p = para(f"{num} {title}", size=14, bold=True, align="both", first=0, before=8, keep=True); use(2.0, atomic=True); (bookmark(p, bm), BM_TEXT.__setitem__(bm, f"{num} {title}")) if bm else None; return p
def h2(num, title, bm=None): p = para(f"{num} {title}", size=12, bold=True, align="both", first=0, before=6, keep=True); use(1.6, atomic=True); (bookmark(p, bm), BM_TEXT.__setitem__(bm, f"{num} {title}")) if bm else None; return p
def fit_picture(path, max_w=6.2, max_h=7.6):
    w, h = Image.open(path).size; sc = min(max_w / (w / 100), max_h / (h / 100)); return w / 100 * sc, h / 100 * sc
def figure(path, caption, max_w=6.2, max_h=7.2):
    w, h = fit_picture(path, max_w, max_h); use(h * 72 / 20.7 + 2.5, atomic=True); doc.add_picture(str(path), width=Inches(w)); p = doc.paragraphs[-1]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True; p.paragraph_format.space_before = Pt(6)
    n = len(FIGS) + 1; c = para(f"Figure {n}: {caption}", size=11, bold=True, align="center", first=0, after=8, spacing=1.15); bookmark(c, f"fig{n}"); BM_TEXT[f"fig{n}"] = f"Figure {n}:"; FIGS.append((n, caption, f"fig{n}")); return n
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill); tcPr.append(sh)
def borders(tbl):
    tblPr = tbl._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"): x = OxmlElement(f"w:{e}"); x.set(qn("w:val"), "single"); x.set(qn("w:sz"), "4"); x.set(qn("w:space"), "0"); x.set(qn("w:color"), "808080"); b.append(x)
    tblPr.append(b)
def table(caption, header, rows, widths=None, size=9.5, align=None, bold_rows=()):
    n = len(TABS) + 1; c = para(f"Table {n}: {caption}", size=11, bold=True, align="center", first=0, before=6, after=4, spacing=1.15, keep=True); bookmark(c, f"tab{n}"); BM_TEXT[f"tab{n}"] = f"Table {n}:"; TABS.append((n, caption, f"tab{n}"))
    est_l = 1.5 + sum(max(1, math.ceil(max(len(str(x)) for x in r) / 28)) * 0.62 for r in [header] + rows); use(min(est_l, LPP - 3), atomic=True)
    if est_l > LPP - 3: use(est_l - (LPP - 3))
    t = doc.add_table(rows=1 + len(rows), cols=len(header)); borders(t); t.alignment = 1
    for j, h in enumerate(header):
        cell = t.rows[0].cells[j]; cell.text = ""; pp = cell.paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER; runs(pp, str(h), size, True); shade(cell, "D9D9D9")
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = t.rows[i + 1].cells[j]; cell.text = ""; pp = cell.paragraphs[0]; pp.alignment = (WD_ALIGN_PARAGRAPH.LEFT if (j == 0 or (align and align[j] == "l")) else WD_ALIGN_PARAGRAPH.CENTER); runs(pp, str(v), size, i in bold_rows)
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths): row.cells[j].width = Inches(w)
    for row in t.rows:
        for cell in row.cells:
            for pp in cell.paragraphs: pp.paragraph_format.line_spacing = 1.0; pp.paragraph_format.space_after = Pt(1); pp.paragraph_format.space_before = Pt(1)
    sp = para("", size=6, first=0, spacing=1.0); use(0.5); return n
def code(lines, size=8):
    for ln in lines:
        ind = len(ln) - len(ln.lstrip()); parts = textwrap.wrap(ln.strip(), width=96 - ind, initial_indent=" " * ind, subsequent_indent=" " * (ind + 6), break_long_words=False) or [" "]
        for q in parts: para(q, size=size, align="left", first=0, spacing=1.0, font="Courier New", left=0.1); use(0.5)
def ref(t): return para(t, size=12, first=0, left=0.45, hanging=0.45, after=4, spacing=1.3)

# ---------- data used in text and tables ----------
models = J("models"); we = J("weather_effect"); tun = J("tuning"); clu = J("clusters"); rules = J("rules"); pv = J("paper_vs_ours"); shap = J("shap"); gains = J("gain_groups"); cause = J("cause_share"); wc = J("weather_class"); prop = J("propagation")
mt = lambda t: sorted([m for m in models if m["task"] == t], key=lambda r: -r["pr_auc"]); best = lambda t: mt(t)[0]
f3 = lambda v: f"{v:.3f}"; p1 = lambda v: f"{100 * v:.1f}%"
SHORT = {"XGBoost (native categoricals, tuned)": "XGBoost (tuned)", "LightGBM (main)": "LightGBM", "XGBoost (tuned, ordinal codes)": "XGBoost (ordinal)", "Decision Tree (CART)": "Decision Tree", "Neural Network (MLP)": "Neural network",
         "Ensemble (LightGBM + XGBoost + MLP)": "Ensemble", "Naive Bayes (Gaussian)": "NB (Gaussian)", "Naive Bayes (Categorical)": "NB (categorical)"}
sh = lambda m: SHORT.get(m, m)
base_rate = 0.1795

# ======================================================================= cover page ===============================================================
for r in P[0].runs:
    if r.text.strip() == "TITLE": r.text = TITLE
    r.font.size = Pt(20)
    if r.text.strip() == "": r.text = ""
pf0 = P[0].paragraph_format; pf0.left_indent = Inches(0.6); pf0.right_indent = Inches(0.6); pf0.first_line_indent = Inches(0); P[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
for q in (P[14], P[15]): q._p.getparent().remove(q._p)
for r in P[12].runs:
    if "CSI4004" in r.text: r.text = "      <Course code and name>"
for r in P[13].runs:
    if "Under the guidance" in r.text: r.text = "Under the guidance of <Faculty name>"
doc.core_properties.title = "Flight Delay and Airport Operations Intelligence"; doc.core_properties.subject = "Project report"

# ======================================================================= remove template body (keep cover, TOC, abstract, lists) ===============================
anchor = P[37]._p; el = anchor; rm = []
while el is not None and el.tag != qn("w:sectPr"): rm.append(el); el = el.getnext()
for e in rm: body.remove(e)

# ======================================================================= CHAPTER 1 ==================================================================
chapter(1, "INTRODUCTION", "ch1")
h1("1.1", "SYSTEM OVERVIEW", "s11")
text("Air travel in the United States moves several million passengers every week, and a flight that arrives late is one of the most visible failures of that system. The U.S. Department of Transportation (DOT) treats a flight as delayed when it arrives at least 15 minutes after its scheduled arrival time [12]. About one flight in five missed that mark in 2015, and the causes were routinely blamed on weather alone, even though airline operations, congestion at the airports and the way a late aircraft carries its delay into its next flights all contribute.")
text("This project builds a complete analytics system around the 2015 DOT on-time performance data. The raw flight records (5,819,079 scheduled domestic flights) and the NOAA GHCN-Daily weather observations [14] are loaded into a star-schema data warehouse. On top of the warehouse the system performs online analytical processing (OLAP), quantifies how much weather really changes delays and cancellations, follows delays along each aircraft's day, and discovers patterns among airports and cancellations. Its central part is a prediction study: will a flight arrive 15 or more minutes late, predicted at two realistic moments. The first is the day before (day-ahead), using only the schedule, the airline, the airports and the weather. The second is one hour before scheduled departure (T-60), when the status of the inbound aircraft is also known.")
text("The study compares eleven models, including gradient-boosted trees, classical learners, a neural network and an ensemble, under a strict time-based evaluation (train on January to August, tune on September, test on November and December). It is benchmarked against a published base paper on the same data [1]. The results are explained with SHAP and presented in an interactive local portal that includes a live flight predictor.")
figure(R / "report_figs" / "prediction_moments.png", "The prediction moments studied: day-ahead, T-60 and (for comparison only) after take-off", max_h=3.0)
h1("1.2", "OBJECTIVE", "s12")
text("The objectives of the project are the following:", first=0)
for t in ["**O1 Data warehouse.** Design a dimensional model and an ETL pipeline that integrates flights, airlines, airports, routes and daily weather.",
          "**O2 OLAP analysis.** Roll-up, drill-down, slice, dice and pivot the data by time, airline and airport to find where and when delays occur.",
          "**O3 Root cause and weather impact.** Separate weather from other causes and measure its effect on delay and cancellation after controlling for airport, month, hour and airline.",
          "**O4 Delay propagation.** Follow each aircraft through its day using its tail number and measure how much of a delay is passed on, absorbed or amplified.",
          "**O5 Predictive mining.** Predict whether a flight arrives 15 or more minutes late at T-60 and day-ahead, compare algorithms fairly, quantify the contribution of weather and explain the predictions.",
          "**O6 Pattern discovery.** Cluster airports by their operating profile and mine association rules for flight cancellations.",
          "**O7 Base-paper comparison.** Compare accuracy, recall and F2 with the base paper [1] and state honestly where the results are higher and where they are not.",
          "**O8 Portal.** Deliver an interactive dashboard with a live predictor so that every result can be inspected and reproduced."]: bullet(t)
h1("1.3", "APPLICATIONS", "s13")
for t in ["**Airline operations.** An alert one hour before departure lets an airline re-time crews, swap aircraft or prepare passenger notifications before the delay is visible at the gate.",
          "**Airport planning.** Day-ahead risk scores and the weather-sensitivity ranking of airports help schedule ground staff, de-icing and gate capacity.",
          "**Schedule design.** The turnaround-buffer analysis shows how long a planned turn must be to absorb an inbound delay, which supports more robust schedules.",
          "**Passenger information.** Early and explainable delay estimates reduce missed connections and improve customer communication.",
          "**Policy and research.** Separating weather, airline and air-system causes gives regulators and researchers a clearer picture of where delay minutes come from."]: bullet(t)
h1("1.4", "LIMITATIONS", "s14")
for t in ["**Weather is observed, not forecast.** Same-day observed weather is used as a stand-in for a forecast, so the measured benefit of weather is an upper bound for a deployed system.",
          "**One year and one test period.** The data cover 2015 only and the models are scored on November and December; other seasons and years are untested.",
          "**October is excluded** from airport, weather and model work because its airport codes in the source file are numeric and could not be mapped to IATA codes.",
          "**Unknown inbound delay.** About 72% of flights have no known inbound arrival delay at T-60 (the inbound aircraft is still airborne, it is the first flight of the day, or it has not left yet), so the T-60 model relies on the inbound aircraft's departure delay and status.",
          "**Associations are not causes.** Cancellation rules and airport clusters describe patterns; clusters are soft and the rules reflect where carriers operate as well as how they behave.",
          "**Comparison with the base paper is approximate** because the splits, airport subsets and class balancing differ; the day-ahead task favours the paper."]: bullet(t)

# ======================================================================= CHAPTER 2 ==================================================================
chapter(2, "SYSTEM ANALYSIS", "ch2")
h1("2.1", "EXISTING SYSTEM", "s21")
text("Flight delay prediction has been studied with many techniques. Gui et al. combine aviation big data and machine learning to predict delays [2], Yazdi et al. train a deep network with the Levenberg-Marquardt algorithm [3], and Belcastro et al. use scalable data mining on large flight and weather collections [4]. The closest work to this project is the base paper by Schösser and Schönberger, who predict whether an arrival is delayed by 15 minutes or more on the same 2015 DOT data using Random Forest, XGBoost and a neural network [1].")
text("The base paper evaluates three feature sets: static schedule and weather features only (set 1), plus the delay of the aircraft's previous flight (set 1+2), and plus the flight's own departure delay (set 1+3). It restricts the data to the 30 busiest airports (2.23 million flights), splits the records randomly into 75% for training and 25% for testing, balances the classes by undersampling the training set, tunes hyperparameters by grid search and particle swarm optimisation, and reports accuracy, recall and the F2 score. Its best model, XGBoost, reaches 69.5% accuracy and 68.2% recall without short-term features, 70.9% and 68.5% with the previous flight's delay, and 89.9% and 83.4% when the departure delay is added.")
text("Existing systems of this kind have several weaknesses that this project addresses:", first=0)
for t in ["A random split puts flights from the same day, and from the same aircraft rotation, in both the training and the test set, which makes the task easier than real forecasting.",
          "Only a subset of airports is used, so the results do not describe the whole network.",
          "Undersampling shifts the operating point towards high recall and low accuracy, and only accuracy, recall and F2 are reported; precision, ranking quality (ROC-AUC, PR-AUC) and uncertainty are not.",
          "Only three algorithms are compared and the effect of weather itself is not isolated.",
          "The work is a prediction study only: there is no data warehouse, no analysis of causes or propagation and no interface for users."]: bullet(t)
h1("2.2", "PROPOSED SYSTEM", "s22")
text("The proposed system is an end-to-end analytics platform. Raw DOT and NOAA files are transformed by an ETL pipeline into a star-schema warehouse in DuckDB [10]. SQL analyses on the warehouse answer the descriptive questions (when, where and why delays occur, and how much weather matters). A leakage-controlled feature table is then built for each prediction moment and eleven models are trained on the early months and scored on the final months. Predictions are explained with SHAP [7], the effect of weather is measured with paired ablations and bootstrap confidence intervals, and the results are compared with the base paper at two operating points. Everything is exposed through a local portal with separate pages for T-60, day-ahead and post-departure, an overall metric-evaluation page and a live predictor.")
h2("2.2.1", "Benefits of Proposed System", "s221")
for t in ["**Realistic evaluation.** A time-based split (train January to August, validate September, test November and December) means no test-month flight is seen during training.",
          "**Whole network.** All airports are used (113 airports with at least 5,000 flights are profiled), not only the 30 busiest.",
          "**Fair comparison.** Eleven models are trained with identical splits, class weights and threshold rules, and ranked by PR-AUC, the metric that stays honest when only 18% of flights are late.",
          "**Isolated weather effect.** Weather features are switched on and off for the tuned models, with paired bootstrap intervals, and the controlled effect on delay and cancellation is measured separately.",
          "**Leakage control.** An assertion fails the run if any feature uses information from after the prediction moment.",
          "**Explainability.** SHAP values, feature-group gains, ablations and subgroup checks show why the model predicts what it does.",
          "**Usability.** A portal lets a reviewer inspect every result and try the model on a real or hypothetical flight."]: bullet(t)

# ======================================================================= CHAPTER 3 ==================================================================
chapter(3, "REQUIREMENT SPECIFICATION", "ch3")
h1("3.1", "HARDWARE REQUIREMENTS", "s31")
text("These are the hardware and software that support the system for smooth operation. The figures below are the configuration on which the project was developed and run.", first=0.5)
for a, b in [("Operating System", "macOS (Darwin 25) on Apple silicon; the code also runs on Linux and Windows with Python 3.9 or later."), ("CPU", "A multi-core processor; the experiments ran on a 10-core machine. The neural network uses the Apple Metal (MPS) GPU when available and falls back to the CPU."),
             ("RAM", "16 GB or more is recommended. DuckDB is capped at 6 GB and the feature table of about five million rows is held in memory when models are trained."), ("Storage", "About 4 GB of free disk space: the DuckDB warehouse (about 0.5 GB) plus a verified backup copy, the cached feature table and model predictions, and the result files."),
             ("Other requirements", "A modern web browser (Google Chrome was used) for the portal. An internet connection is only needed to install the Python packages; the portal itself runs offline on localhost.")]:
    p = para(a, size=12, bold=True, first=0, left=0.3, spacing=1.5, after=0, keep=True); use(1.1); text(b, first=0.5, left=0.3)
h1("3.2", "SOFTWARE REQUIREMENTS", "s32")
text("The software description follows the different tools and libraries used for building, analysing and presenting the system.", first=0.5)
items = [("1 Python 3.9 and virtual environment", "All analysis and modelling code is written in Python 3.9 inside an isolated virtual environment (.venv) with the package versions pinned in requirements.txt, so the results can be reproduced."),
         ("2 DuckDB 1.4", "DuckDB is an embeddable analytical (OLAP) database [10]. The star-schema warehouse is stored in a single file and queried with SQL, including window functions used for the aircraft-chain analysis. Every analysis script opens it read-only."),
         ("3 pandas, NumPy and PyArrow", "pandas and NumPy handle tabular data and numerical work; PyArrow reads and writes Parquet files used to cache the feature table."),
         ("4 scikit-learn 1.6", "scikit-learn [8] provides Random Forest, Decision Tree, Logistic Regression, KNN, Naive Bayes, the preprocessing pipelines, K-Means, hierarchical clustering, PCA and the evaluation metrics."),
         ("5 LightGBM 4.6 and XGBoost 2.1", "Gradient-boosted decision trees [5][6], the strongest model family in this study. LightGBM is the main model and also serves the live predictor."),
         ("6 PyTorch 2.8", "PyTorch [11] is used for the embedding multilayer perceptron (MLP). Training runs on the Metal GPU when it is available."),
         ("7 SHAP, statsmodels, SciPy and mlxtend", "SHAP [7] explains predictions; statsmodels fits the logistic regressions with clustered standard errors; SciPy provides rank statistics; mlxtend provides FP-Growth for association rules."),
         ("8 matplotlib and seaborn", "Used for the static charts in the report and for the diagrams in Chapter 4."),
         ("9 d3.js and plain web technologies", "The portal is written in HTML, CSS and JavaScript with d3 v7 for charts. A small Python server built on the standard library serves the files and a JSON API."),
         ("10 python-docx and headless Chrome", "python-docx builds this report from the university template; headless Chrome captured the portal screenshots in Appendix 2.")]
for a, b in items: p = para(a, size=12, bold=True, first=0, spacing=1.5, keep=True); use(1.1); text(b, first=0.5)

# ======================================================================= CHAPTER 4 ==================================================================
chapter(4, "SYSTEM DESIGN SPECIFICATION", "ch4")
h1("4.1", "SYSTEM ARCHITECTURE", "s41")
text("The system is organised as the layered pipeline shown in Figure 2. Source files enter an ETL layer that writes the warehouse; analysis modules read the warehouse, a feature layer prepares model inputs, the modelling and evaluation layer produces the results, and the portal and report present them.")
figure(R / "report_figs" / "architecture.png", "System architecture", max_h=5.6)
for a, b in [("Source layer.", "The DOT on-time file (flights.csv, 5,819,079 rows), the airline and airport lookup files, the GHCN-Daily station list and the 37-million-row daily weather file."),
             ("ETL layer.", "A Python pipeline (etl_pipeline.py) extracts the files, streams the weather file in chunks and keeps only quality-checked precipitation, snow, snow depth, temperature and wind readings, maps every airport to up to three suitable stations within 50 km, builds the dimensions and the fact table with surrogate keys, and writes CSV extracts that are loaded into DuckDB."),
             ("Warehouse layer.", "One fact table and five dimensions (Section 4.3) plus two derived tables that hold the aircraft-chain features."),
             ("Analysis layer.", "Scripts 00 to 06 produce the audit, the OLAP exports, the root-cause and weather analysis, the delay-propagation study, the airport clusters and the cancellation rules."),
             ("Feature and model layer.", "A single SQL query builds the T-60 feature table; scenarios select feature subsets. Models are trained on the early months and their predictions are cached so that any result can be re-computed without retraining."),
             ("Presentation layer.", "A local web portal (Python server with a JSON API and a d3 front end), this report and the screenshots.")]: bullet(f"**{a}** {b}")
h1("4.2", "DETAILED DESIGN", "s42")
h2("4.2.1", "Data preparation and quality checks", "s421")
text("The audit step (00_audit.py) verified that the warehouse has no duplicate keys, no orphan foreign keys, no flight that is both cancelled and diverted, and that every cancelled flight has a reason. It also found the October airport-code problem: the source uses numeric BTS airport identifiers for October, which the ETL mapped to the Unknown airport member by design, so October is kept for time-only OLAP and excluded from airport-level, weather and modelling work. Delayed is defined as an arrival delay of at least 15 minutes, matching the DOT rule and the base paper; flown flights exclude cancelled and diverted flights.")
h2("4.2.2", "Weather classification", "s422")
text("Each airport-day is classified from the observed daily measurements at the nearest station. Severe weather is snowfall above 25 mm, precipitation above 25 mm or average wind above 13.4 m/s; moderate weather is any snowfall, precipitation above 5 mm or wind above 8.9 m/s; everything else is clear. Missing measurements are treated as zero and airports with no station within the search radius are labelled unknown.")
h2("4.2.3", "Delay propagation along an aircraft's day", "s423")
text("For every flight with a tail number the legs of the aircraft are ordered by scheduled departure within each scheduled date using SQL window functions (LAG over all scheduled legs, so a cancelled leg breaks the chain instead of being skipped). Local clock times are converted to an absolute minute axis, which handles time zones and midnight rollover. A previous leg is accepted only if it belongs to the same aircraft and date, was flown, and arrived at this flight's origin; 98.9% of flown previous legs end at the next flight's origin, which confirms the chains are physically consistent. The analysis set has 3,917,656 flights; the subset whose previous arrival is known before the scheduled departure (3,693,636 flights) feeds the prediction features.")
h2("4.2.4", "Prediction moments, features and leakage control", "s424")
text("The T-60 feature table is built by one SQL query over the star schema and holds 5,231,130 flown flights from non-October months. Every column is known at prediction time: schedule and calendar attributes, airline, airports, route, observed weather at both ends (used as a stand-in for a forecast), the status of the inbound aircraft at T-60 (landed, airborne, not yet departed or no previous leg) with its departure and arrival delay where known, the planned turnaround, the leg number of the day, and historical delay rates of the route, airline and origin computed only from earlier training months. Day-ahead drops the six inbound-aircraft features; the post-departure scenario adds the flight's own departure delay and is used only to compare with the base paper. An assertion checks every feature against a forbidden list and verifies in the database that no inbound value uses an event later than T-60.")
table("Feature sets of the three scenarios", ["Scenario", "Features", "Information available"], [["Day-ahead", "37", "Schedule, calendar, airline, airports, route, observed weather at both ends, historical delay rates"], ["T-60", "43", "Day-ahead features + inbound aircraft status, departure and arrival delay, planned turnaround, leg number"], ["Post-departure", "44", "T-60 features + the flight's own departure delay (known only after take-off)"]], widths=[1.2, 0.8, 4.4], align=["l", "c", "l"])
h2("4.2.5", "Models", "s425")
text("Eleven models are compared on identical splits. The training rows are Jan to Aug (3,836,893 flights), the validation rows are September (462,153) and the test rows are November and December (932,084; 17.95% delayed). Classes are weighted by w = n / (2 n_c) in every model that supports weights, and a single decision threshold per model is chosen to maximise F1 on September and applied unchanged to the test flights.")
table("Models and settings", ["Model", "Training rows", "Main settings"], [["LightGBM", "3.84 M", "63 leaves, learning rate 0.05, min child samples 100, feature fraction 0.6, native categoricals, early stopping"], ["XGBoost (tuned)", "3.84 M", "depth 8, learning rate 0.05, min child weight 50, column sample 0.6, native categoricals"],
        ["Random Forest", "1 M sample", "150 trees, minimum leaf 50, sqrt features"], ["Decision Tree", "3.84 M", "CART, depth 12 chosen on validation"], ["Logistic Regression", "3.84 M", "one-hot (rare levels merged), scaled numerics, C = 1"], ["KNN", "200 k sample", "k = 50, numeric features"],
        ["Naive Bayes", "3.84 M", "Gaussian and categorical variants, balanced priors"], ["Neural network", "3.84 M", "embeddings + 512-256-128 MLP, batch norm, dropout 0.25, class-weighted loss, AdamW, early stopping"], ["Ensemble", "-", "mean of rank-normalised scores of LightGBM, XGBoost and the MLP"]], widths=[1.4, 1.1, 3.9], align=["l", "c", "l"], size=9)
h2("4.2.6", "Evaluation protocol and tuning", "s426")
text("PR-AUC (average precision) is the primary metric because it remains meaningful when only 18% of flights are late; ROC-AUC, precision, recall, F1, F2 and accuracy are also reported, with accuracy marked as misleading (always predicting on-time already scores about 82%). Confidence intervals come from a paired Poisson bootstrap over test flights, and the boosted models were run with three seeds. Hyperparameters were first chosen by grid search on a 1-million-row sample against validation PR-AUC; as in the base paper, particle swarm optimisation (8 particles, 5 iterations, written for this project) was then run for XGBoost, LightGBM and the neural network.")
h2("4.2.7", "Explainability, patterns and the portal", "s427")
text("SHAP (exact TreeSHAP for LightGBM) gives per-flight and global explanations; feature-group gain, a depth-4 decision tree and subgroup scores complement them. Airports with at least 5,000 flights are clustered with K-Means (k chosen by silhouette among 3 to 6 and cross-checked with Ward clustering) on ten features, and cancellation association rules are mined exhaustively over conjunctions of up to three conditions with support, confidence and lift. The portal consists of a Python server that loads the saved T-60 LightGBM model and exposes /api/options, /api/predict and /api/sample, and a plain JavaScript front end (sidebar, tabs, sortable tables and d3 charts) with separate pages for T-60, day-ahead, post-departure, overall metric evaluation, data exploration and an About page.")
h1("4.3", "DATABASE DESIGN", "s43")
text("The warehouse is a star schema with one fact table at the grain of a scheduled flight leg and five dimensions; the airport and weather dimensions are role-played for the origin and the destination (Figure 3 and Table 3).")
figure(R / "report_figs" / "star_schema.png", "Star schema of the flight data warehouse", max_h=6.2)
table("Warehouse tables", ["Table", "Grain / content", "Key columns"], [["fact_flight", "One row per scheduled flight leg (5,819,079 rows)", "flight_key; foreign keys to date, airline, route, origin and destination airport and weather; tail_number, flight_number; times; delay and cause minutes; cancelled and diverted flags"],
        ["dim_date", "One row per calendar day", "date_key, full_date, month, quarter, day of week, season, weekend and holiday flags"], ["dim_airline", "One row per carrier", "airline_key, IATA code, name"], ["dim_airport", "One row per airport", "airport_key, IATA code, name, city, state, latitude, longitude"],
        ["dim_route", "One row per origin-destination pair", "route_key, origin, destination, distance band, interstate flag"], ["dim_weather", "One row per airport per day", "weather_key, precipitation, snowfall, snow depth, tmax, tmin, wind, station and distance"],
        ["propagation_features", "Derived: aircraft-chain features (analysis set)", "previous leg arrival and departure delay, planned turn, leg number"], ["propagation_t60", "Derived: inbound status at T-60", "inbound status, delays known at T-60, absolute-time checks"]], widths=[1.4, 2.2, 2.9], align=["l", "l", "l"], size=9)

# ======================================================================= CHAPTER 5 ==================================================================
chapter(5, "RESULT AND CONCLUSION", "ch5")
h1("5.1", "RESULT", "s51")
h2("5.1.1", "When and where flights are late (OLAP)", "s511")
text("Delay builds through the day: the on-time share falls from 92.7% for 05:00 departures to 73.7% at 20:00, and evening departures are 2.4 times as likely to leave 15 or more minutes late as morning departures (27.0% against 11.3%). The worst months are February (76.65% on time, 4.78% cancelled) and June (76.52%), the best is September (87.00%). Saturday is the best weekday (84.01% on time) and Monday has the highest cancellation rate (2.43%). Hawaiian Airlines (88.67%), Alaska (86.96%) and Delta (86.44%) are the most punctual large carriers; Spirit (70.29%) and Frontier (73.84%) are the least, and American Eagle cancels 5.1% of its flights, three times the network rate.")
figure(R / "olap" / "by_dep_hour.png", "On-time share by scheduled departure hour", max_h=3.6)
h2("5.1.2", "Weather and the causes of delay", "s512")
sv = next(r for r in wc if r["wx_class"] == "severe"); mo = next(r for r in wc if r["wx_class"] == "moderate"); cl = next(r for r in wc if r["wx_class"] == "clear")
text(f"Severe origin weather raises the share of delayed flights from {cl['delay_rate_pct']:.1f}% on clear days to {sv['delay_rate_pct']:.1f}% (moderate: {mo['delay_rate_pct']:.1f}%) and the cancellation rate from {cl['cancel_rate_pct']:.1f}% to {sv['cancel_rate_pct']:.1f}%. Controlling for airport and month leaves the gap unchanged (+23.4 percentage points, about +26 minutes of arrival delay), and adjusting for airport, month, hour and airline gives odds ratios of 3.60 for a delay and 7.70 for a cancellation in severe weather. Weather causes 54.8% of all cancellations (47,691 of 87,076). Among delayed flights, late aircraft account for 39.8% of delay minutes, airline operations for 32.2%, the air system for 22.9% and weather for only 4.9%; DOT's weather code captures only about a fifth of the extra delay on severe-weather days because air-system and late-aircraft delays rise as well.")
table("Delay-cause shares for flights 15 or more minutes late", ["Cause", "Delay minutes", "Share"], [[c["cause"], f"{c['delay_minutes']:,.0f}", f"{100 * c['delay_minutes'] / sum(x['delay_minutes'] for x in cause):.1f}%"] for c in cause], widths=[2.4, 2.0, 1.2])
figure(R / "root_cause_weather" / "2c_raw_vs_controlled.png", "Weather effect on delay and cancellation, raw and controlled for airport and month", max_h=4.2)
h2("5.1.3", "Delay propagation", "s513")
text("Of all departure delays of 15 minutes or more in the analysis set, 60.8% follow an inbound flight that arrived at least 15 minutes late. When the inbound aircraft is 30 to 59 minutes late, 83.2% of the next departures leave 15 or more minutes late, against 16.5% when the inbound arrival was only 0 to 14 minutes late. The mean departure delay rises from 5.0 minutes on an aircraft's first leg to 16.4 minutes on its seventh. A planned turnaround of 90 minutes or more absorbs about three quarters of delays under an hour, while a turn under 45 minutes absorbs nothing. SkyWest and Hawaiian absorb the most inbound delay, while United, JetBlue and Spirit amplify it.")
figure(R / "propagation" / "3A_propagation_curve.png", "Departure delay against the inbound aircraft's arrival delay", max_h=3.8)
h2("5.1.4", "Prediction at T-60 and day-ahead", "s514")
text(f"At T-60 the ensemble of LightGBM, XGBoost and the neural network ranks flights best (PR-AUC {f3(best('T-60')['pr_auc'])}, ROC-AUC {f3(best('T-60')['roc_auc'])}, a lift of {best('T-60')['pr_auc'] / base_rate:.1f} times over the 18% chance level), closely followed by tuned XGBoost and LightGBM; the three are statistically close and the ensemble's lead is small. The neural network ({f3(next(m for m in mt('T-60') if m['model'].startswith('Neural'))['pr_auc'])}) trails the boosted trees, as is typical for tabular data, but beats Random Forest. Naive Bayes and KNN are weakest. Day-ahead is much harder (best PR-AUC {f3(best('Day-ahead')['pr_auc'])}, {best('Day-ahead')['pr_auc'] / base_rate:.1f} times chance) because nothing is known about the aircraft's earlier legs.")
hdr = ["Model", "PR-AUC", "ROC-AUC", "Accuracy", "Precision", "Recall", "F1", "F2"]
rowf = lambda m: [sh(m["model"]), f3(m["pr_auc"]), f3(m["roc_auc"]), p1(m["accuracy"]), p1(m["precision"]), p1(m["recall"]), f3(m["f1"]), f3(m["f2"])]
table("T-60 results on the November and December test flights (F1-tuned threshold)", hdr, [rowf(m) for m in mt("T-60")], widths=[1.6, 0.7, 0.75, 0.75, 0.75, 0.65, 0.6, 0.6], size=9, bold_rows=(0,))
table("Day-ahead results on the November and December test flights (F1-tuned threshold)", hdr, [rowf(m) for m in mt("Day-ahead") if m["model"] != "Majority dummy"], widths=[1.6, 0.7, 0.75, 0.75, 0.75, 0.65, 0.6, 0.6], size=9, bold_rows=(0,))
figure(R / "model_comparison" / "model_comparison_t60.png", "T-60 model comparison by PR-AUC with bootstrap 95% intervals (boosted and classical models; the neural network and ensemble are in Table 5)", max_h=4.0)
figure(R / "model_comparison" / "pr_curves_t60.png", "Precision-recall curves at T-60", max_h=4.0)
h2("5.1.5", "Post-departure scenario", "s515")
text(f"When the flight's own departure delay is added, ranking quality rises sharply (XGBoost PR-AUC {f3(best('Post-departure')['pr_auc'])}, ROC-AUC {f3(best('Post-departure')['roc_auc'])}, {best('Post-departure')['pr_auc'] / base_rate:.1f} times chance), but this number is known only after take-off, so the scenario exists to compare with the base paper and to show how much a single piece of information is worth.")
table("Post-departure results (F1-tuned threshold)", hdr, [rowf(m) for m in mt("Post-departure")], widths=[1.6, 0.7, 0.75, 0.75, 0.75, 0.65, 0.6, 0.6], size=9)
h2("5.1.6", "Effect of weather on prediction", "s516")
text("Switching the weather features off and on for the tuned models shows how much weather contributes. The effect is large when nothing else is known and smaller, but still clearly non-zero, when the inbound aircraft is known; all four intervals exclude zero. Weather also accounts for 11.7% of LightGBM's gain, behind the inbound aircraft's status and delay (42.3%) and airline, airport and route (24.3%). The weather used is observed same-day weather, so these gains are upper bounds for a real forecast.")
table("PR-AUC with and without weather features (tuned models)", ["Scenario", "Model", "Without", "With", "Gain (95% interval)"], [[w["task"], w["model"].replace(" (PSO-tuned)", ""), f3(w["pr_auc_off"]), f3(w["pr_auc_on"]), f"+{w['delta_pr_auc']:.3f} [{w['ci_lo']:.3f}, {w['ci_hi']:.3f}]"] for w in we], widths=[1.1, 1.2, 0.9, 0.9, 2.0], align=["l", "l", "c", "c", "c"])
figure(R / "model_comparison" / "ablation_lightgbm.png", "LightGBM ablation: effect of weather, propagation and history features", max_h=3.6)
h2("5.1.7", "Explainability", "s517")
text("The mean absolute SHAP values show that the most influential features at T-60 are the inbound aircraft's departure delay, the time of year, destination precipitation and the inbound status. A late inbound aircraft pushes the risk up sharply, while a long planned turnaround and early departure hours lower it.")
figure(R / "explain" / "shap_summary.png", "SHAP summary of the T-60 model", max_h=4.6)
h2("5.1.8", "Hyperparameter tuning", "s518")
text("Particle swarm tuning, run on a 600,000-row training sample, improved T-60 LightGBM by 0.004 PR-AUC and changed the other models by less than that or made them slightly worse, because settings chosen on a sample did not carry over to full-data training. The earlier grid-searched configurations therefore remain the reference results.")
table("Particle swarm tuning results", ["Scenario", "Model", "Tuned values", "Before", "After"], [[t["task"], t["model"], t["values"], f3(t["pr_auc_before"]), f3(t["pr_auc_tuned"])] for t in tun], widths=[0.9, 1.0, 3.4, 0.6, 0.6], align=["l", "l", "l", "c", "c"], size=8.5)
h2("5.1.9", "Airport clusters and cancellation rules", "s519")
text("K-Means on 113 airports gives three groups: large hubs with the longest delays and the greatest sensitivity to adverse weather, mid-size airports that are the most punctual, and smaller airports with the highest cancellation rate and weather exposure. The clusters are soft (silhouette 0.24) but stable under resampling (mean adjusted Rand index 0.88) and agree with Ward clustering (0.82). Exhaustive rule mining over 5.30 million flights (87,076 cancelled, base rate 1.64%) shows that cancellations are concentrated where severe weather, snow, high wind and freezing temperatures combine, with lifts of 20 to 28; the regional carriers MQ and EV recur in the strongest rules. These are associations, not causes.")
table("Airport clusters", ["Cluster", "Airports", "Delayed", "Cancelled", "Weather gap (pp)", "Examples"], [[c["name"], c["airports"], f"{c['delayed_pct']:.1f}%", f"{c['cancel_rate_pct']:.2f}%", f"{c['weather_gap_pp']:.1f}", c["example_airports"].split(", ")[0:5] and ", ".join(c["example_airports"].split(", ")[0:5])] for c in clu], widths=[1.6, 0.7, 0.7, 0.8, 0.9, 1.8], align=["l", "c", "c", "c", "c", "l"], size=9)
table("Strongest non-redundant cancellation rules", ["If the flight has", "Flights", "Cancelled", "Share", "Lift"], [[r["antecedent"], f"{r['antecedent_flights']:,}", f"{r['cancelled']:,.0f}", f"{100 * r['confidence']:.1f}%", f"{r['lift']:.1f}x"] for r in rules[:6]], widths=[3.4, 0.7, 0.8, 0.7, 0.6], align=["l", "c", "c", "c", "c"], size=9)
figure(R / "clustering" / "airport_clusters_pca.png", "Airport clusters (PCA projection)", max_h=4.0)
figure(R / "association_rules" / "top_cancellation_rules.png", "Cancellation rules with the highest lift", max_h=4.0)
h2("5.1.10", "Comparison with the base paper", "s5110")
text(f"Both studies use the same 2015 data and the same 15-minute rule, and both find tree ensembles ahead of the neural network and both find that the inbound aircraft's delay helps. The comparison is approximate because the base paper uses a random split on 30 airports with undersampled training, so Table {len(TABS) + 1} shows our models at their own F1-tuned threshold and Table {len(TABS) + 2} at a lowered threshold whose recall is close to the paper's. At our own threshold we exceed the paper's accuracy in every scenario (by 0.3 points day-ahead, 14.3 at T-60 and 3.2 post-departure for XGBoost) but have lower recall and F2, as expected when the threshold favours precision. At matched recall the T-60 models are clearly higher (ensemble 79.2% accuracy and 61.0% F2 against 70.9% and 58.1%), the post-departure models are level on F2 (80.0% against 80.3%), and the paper is ahead on the day-ahead task.")
def pcmp(scn, th):
    px = next(r for r in pv if r["scenario"] == scn and r["system"] == "Paper XGBoost"); rows = [["Paper XGBoost", f"{px['accuracy']:.1f}", "-", f"{px['recall']:.1f}", f"{px['f2']:.1f}", "-", "-", "-"]]
    for r in [r for r in pv if r["scenario"] == scn and r["source"] == "ours" and r["threshold"] == th]:
        rows.append([r["system"].replace("Ours ", "Ours: ").replace(" (F1 threshold)", "").replace(" (recall matched)", ""), f"{r['accuracy']:.1f}", f"{r['precision']:.1f}", f"{r['recall']:.1f}", f"{r['f2']:.1f}", f"{r['accuracy'] - px['accuracy']:+.1f}", f"{r['recall'] - px['recall']:+.1f}", f"{r['f2'] - px['f2']:+.1f}"])
    return rows
ph = ["System", "Acc.", "Prec.", "Recall", "F2", "Δ Acc.", "Δ Rec.", "Δ F2"]
rows11 = []; rows12 = []
for scn in ["Day-ahead", "T-60", "Post-departure"]:
    rows11.append([f"**{scn}**", "", "", "", "", "", "", ""]); rows11 += pcmp(scn, "F1-tuned"); rows12.append([f"**{scn}**", "", "", "", "", "", "", ""]); rows12 += pcmp(scn, "Recall matched to paper")[0:1] + pcmp(scn, "Recall matched to paper")[1:]
def strip(rows): return [[c.replace("**", "") for c in r] for r in rows]
table("Base paper versus ours at our F1-tuned threshold (percent; differences in points)", ph, strip(rows11), widths=[2.1, 0.6, 0.6, 0.65, 0.55, 0.6, 0.6, 0.55], size=8.5, bold_rows=tuple(i for i, r in enumerate(rows11) if r[0].startswith("**")))
table("Base paper versus ours with the threshold lowered to match the paper's recall", ph, strip(rows12), widths=[2.1, 0.6, 0.6, 0.65, 0.55, 0.6, 0.6, 0.55], size=8.5, bold_rows=tuple(i for i, r in enumerate(rows12) if r[0].startswith("**")))
figure(R / "paper_comparison.png", "Accuracy and F2 against the base paper in the three scenarios", max_h=3.4)
h2("5.1.11", "The portal", "s5111")
text("The local portal (Appendix 2) has a page for each prediction moment, an overall metric-evaluation page with confidence intervals, a model-by-metric heatmap, precision-recall, ROC and lift curves and a threshold sweep, a data-exploration page, and a live predictor. The predictor returns a probability, the alert verdict relative to the September-tuned threshold and the six features that pushed the estimate, and can load a real November or December flight to compare the prediction with what actually happened.")
h1("5.2", "CONCLUSION", "s52")
text("This project built a data warehouse from 5.8 million 2015 US flights and NOAA weather and used it to answer two questions. First, weather matters: severe weather at the origin multiplies the delay rate by 2.4 (17.0% to 40.4%) and the cancellation rate by about 8, and the effect survives controls for airport, month, hour and airline. Second, early prediction is feasible and its quality depends mainly on what is known: the best PR-AUC rises from 0.36 day-ahead to 0.64 at T-60 and 0.88 once the departure delay is known, while weather features add 0.10 and 0.04 PR-AUC at the two operational moments. A late inbound aircraft is the strongest single signal, and planned turnaround buffers of 90 minutes or more absorb most delays under an hour.")
text("Against the base paper, the results are higher where the comparison is like for like: the T-60 models exceed the paper's accuracy and F2 at the same recall, and at our own threshold accuracy is higher in every scenario. The paper remains ahead on the day-ahead task and on recall at our F1 threshold, and the comparison is approximate because of different splits, airport subsets and class balancing. The evaluation here is stricter (time-based split, all airports, confidence intervals, PR-AUC and ROC-AUC reported) and the study is broader (warehouse, root-cause analysis, propagation, clustering, association rules and an interactive portal).")
text("Future work includes using real weather forecasts instead of observed weather, extending the data to several years and to international flights, adding crew and airport-capacity information, calibrating the probabilities for operational use, deploying the predictor as a service, and training a day-ahead model that can be queried in the portal.", before=2)

# ======================================================================= CHAPTER 6 ==================================================================
chapter(6, "APPENDICES", "ch6")
h1("6.1", "APPENDIX 1 - SAMPLE SOURCE CODE", "s61")
text("The listings below are short excerpts of the project code (full scripts are in the mining/ folder).", first=0)
h2("A. Leakage assertion (mining/ml_common.py)", "", None) if False else None
para("A. Leakage assertion (excerpt)", size=12, bold=True, first=0, keep=True); use(1.3)
code(["def assert_no_leakage(features, verbose=True, allow_rebuild=True):", "    feats = list(features)", "    bad = [c for c in feats if c in FORBIDDEN_EXACT or (FORBIDDEN_PATTERN.search(c) and c not in ALLOWED_PREV)]",
      "    assert not bad, f\"LEAKAGE: forbidden columns in feature matrix: {bad}\"", "    unknown = [c for c in feats if c not in ALL_FEATURES]", "    assert not unknown, f\"Features outside the approved whitelist: {unknown}\"",
      "    # database check: no inbound-aircraft value may use an event later than T-60", "    r = con.execute(\"SELECT sum((prev_leg_arr_delay IS NOT NULL AND (prev_arr_abs > t60_abs", "                       OR prev_leg_status <> 'landed'))::int), ... FROM propagation_t60\").fetchone()", "    assert all(v == 0 for v in r), f\"LEAKAGE: inbound feature uses an event after T-60: {r}\""])
para("B. Particle swarm optimiser (mining/pso.py)", size=12, bold=True, first=0, keep=True, before=8); use(1.3)
code(["def pso(objective, space, n_particles=8, n_iter=5, seed=42, w=0.7, c1=1.5, c2=1.5, log_path=None):", "    rng = np.random.default_rng(seed); d = len(space)", "    pos = rng.random((n_particles, d)); vel = rng.normal(0, 0.1, (n_particles, d))",
      "    pbest, pscore = pos.copy(), np.full(n_particles, -np.inf); gbest, gscore = pos[0].copy(), -np.inf", "    for it in range(n_iter + 1):", "        for i in range(n_particles):", "            p = decode(pos[i], space); s = float(objective(p))",
      "            if s > pscore[i]: pscore[i], pbest[i] = s, pos[i].copy()", "            if s > gscore: gscore, gbest = s, pos[i].copy()", "        r1, r2 = rng.random((n_particles, d)), rng.random((n_particles, d))",
      "        vel = w * vel + c1 * r1 * (pbest - pos) + c2 * r2 * (gbest - pos); pos = np.clip(pos + vel, 0, 1)", "    return decode(gbest, space), gscore"])
para("C. Aircraft-chain window function (excerpt of the propagation SQL)", size=12, bold=True, first=0, keep=True, before=8); use(1.3)
code(["SELECT *, row_number() OVER p AS leg_number,", "       lag(flown) OVER p AS prev_flown, lag(d) OVER p AS prev_dest, lag(arr_abs) OVER p AS prev_arr_abs,", "       lag(dd) OVER p AS prev_dep_delay, lag(ad) OVER p AS prev_arr_delay", "FROM c",
      "WINDOW p AS (PARTITION BY tail_number, date_key ORDER BY sdep, flight_number, flight_key)"])
para("D. Prediction API (excerpt of dashboard/server.py)", size=12, bold=True, first=0, keep=True, before=8); use(1.3)
code(["def score(X):", "    X = X[FEATS].copy()", "    for c in CAT: X[c] = pd.Categorical(X[c])", "    with lock: prob = float(bst.predict(X)[0]); con = bst.predict(X, pred_contrib=True)[0]",
      "    top = np.argsort(-np.abs(con[:-1]))[:6]                     # six largest SHAP contributions", "    return dict(probability=prob, threshold=THR, predicted_delayed=prob >= THR, base_rate=BASE_RATE, lift=prob / BASE_RATE,", "                drivers=[...])"])
h1("6.2", "APPENDIX 2 - SCREEN SHOTS / OUTPUTS", "s62")
text("The following screenshots show the portal (light theme) captured with headless Chrome. The predictor screenshot shows a real flight from the November and December test period together with what actually happened.", first=0)
for f, c in [("overview.png", "Portal: overview page"), ("t60_results.png", "Portal: T-60 performance page"), ("t60_predict.png", "Portal: live predictor with a real flight"), ("dayahead_results.png", "Portal: day-ahead performance page"), ("postdep_results.png", "Portal: post-departure performance page"),
             ("evaluation_scenarios.png", "Portal: overall metric evaluation, scenarios"), ("evaluation_models.png", "Portal: overall metric evaluation, models"), ("evaluation_curves.png", "Portal: precision-recall, ROC and lift curves and the threshold sweep"), ("evaluation_paper.png", "Portal: scorecard against the base paper"), ("explore_airports.png", "Portal: airports and cancellations")]:
    figure(SC / f, c, max_h=8.0, max_w=6.0)

# ======================================================================= CHAPTER 7 ==================================================================
chapter(7, "REFERENCES", "ch7")
h1("7.1", "LIST OF JOURNALS", "s71")
J_REFS = ["[1] Schösser D., Schönberger J. On the performance of machine learning based flight delay prediction - investigating the impact of short-term features. Promet - Traffic&Transportation. 2022;34(6):825-838.",
          "[2] Gui G., Liu F., Sun J., Yang J., Zhou Z., Zhao D. Flight delay prediction based on aviation big data and machine learning. IEEE Transactions on Vehicular Technology. 2020;69(1):140-150.",
          "[3] Yazdi M.F., Kamel S.R., Chabok S.J.M., Kheirabadi M. Flight delay prediction based on deep learning and Levenberg-Marquart algorithm. Journal of Big Data. 2020;7:106.",
          "[4] Belcastro L., Marozzo F., Talia D., Trunfio P. Using scalable data mining for predicting flight delays. ACM Transactions on Intelligent Systems and Technology. 2016;8(1):1-20.",
          "[5] Ke G., Meng Q., Finley T., et al. LightGBM: a highly efficient gradient boosting decision tree. Advances in Neural Information Processing Systems 30. 2017.",
          "[6] Chen T., Guestrin C. XGBoost: a scalable tree boosting system. Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining. 2016:785-794.",
          "[7] Lundberg S.M., Lee S.-I. A unified approach to interpreting model predictions. Advances in Neural Information Processing Systems 30. 2017.",
          "[8] Pedregosa F., Varoquaux G., Gramfort A., et al. Scikit-learn: machine learning in Python. Journal of Machine Learning Research. 2011;12:2825-2830.",
          "[9] Menne M.J., Durre I., Vose R.S., Gleason B.E., Houston T.G. An overview of the Global Historical Climatology Network-Daily database. Journal of Atmospheric and Oceanic Technology. 2012;29(7):897-910.",
          "[10] Raasveldt M., Mühleisen H. DuckDB: an embeddable analytical database. Proceedings of the 2019 International Conference on Management of Data (SIGMOD). 2019:1981-1984.",
          "[11] Paszke A., Gross S., Massa F., et al. PyTorch: an imperative style, high-performance deep learning library. Advances in Neural Information Processing Systems 32. 2019."]
for t in J_REFS: ref(t); use(lines_of(t, 75) + 0.3)
h1("7.2", "LIST OF WEBSITES (URLs)", "s72")
for t in ["[12] Bureau of Transportation Statistics. Airline on-time performance and causes of flight delays. https://www.bts.gov/topics/airlines-and-airports/airline-time-performance-and-causes-flight-delays",
          "[13] U.S. Department of Transportation. 2015 Flight Delays and Cancellations (dataset). https://www.kaggle.com/datasets/usdot/flight-delays",
          "[14] NOAA National Centers for Environmental Information. Global Historical Climatology Network - Daily (GHCN-Daily). https://www.ncei.noaa.gov/products/land-based-station/global-historical-climatology-network-daily",
          "[15] DuckDB documentation. https://duckdb.org/docs/", "[16] D3.js - Data-Driven Documents. https://d3js.org", "[17] SHAP documentation. https://shap.readthedocs.io"]: ref(t); use(lines_of(t, 75) + 0.3)
h1("7.3", "LIST OF BOOKS", "s73")
for t in ["[18] Kimball R., Ross M. The Data Warehouse Toolkit: The Definitive Guide to Dimensional Modeling. 3rd ed. Wiley; 2013.", "[19] Han J., Kamber M., Pei J. Data Mining: Concepts and Techniques. 3rd ed. Morgan Kaufmann; 2011.",
          "[20] Hastie T., Tibshirani R., Friedman J. The Elements of Statistical Learning. 2nd ed. Springer; 2009.", "[21] McKinney W. Python for Data Analysis. 3rd ed. O'Reilly Media; 2022."]: ref(t); use(lines_of(t, 75) + 0.3)

# ======================================================================= front matter (abstract, lists, TOC) =============================================
ABS = ["Flight delays cost passengers and airlines time and money, and their causes are usually blamed on weather alone. This project builds a data warehouse from 5.8 million 2015 US flights and NOAA daily weather and uses it to measure how much weather changes delays and cancellations and to predict, at two realistic moments, whether a flight will arrive 15 or more minutes late. A delay is defined by the DOT rule, the same as in the base paper.",
       "A star-schema warehouse in DuckDB supports OLAP, root-cause analysis and a study of how delays propagate along an aircraft's day. Eleven models are compared on a strict time-based split (train January to August, validate September, test November to December): LightGBM, XGBoost, Random Forest, Decision Tree, Logistic Regression, KNN, two Naive Bayes variants, a neural network and an ensemble, tuned by grid search and particle swarm optimisation and explained with SHAP.",
       "Severe weather raises the share of delayed flights from 17.0% to 40.4% and the cancellation rate from 1.1% to 9.0%. At T-60 the best model reaches PR-AUC 0.645 (3.6 times the chance level) and day-ahead 0.360; weather features add 0.04 and 0.10 PR-AUC respectively. Compared with the base paper, the T-60 models are higher in accuracy and F2 at the same recall and in accuracy in every scenario at our own threshold, while the paper is ahead on the day-ahead task. The results are delivered in an interactive portal with a live flight predictor."]
for i, t in enumerate(ABS):
    p = P[23 + i]; pf = p.paragraph_format; pf.line_spacing = 1.5; pf.first_line_indent = Inches(0.5); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; runs(p, t, 12)
bookmark(P[22], "abstract"); bookmark(P[28], "lof"); bookmark(P[31], "loa"); bookmark(P[34], "lot")

def fill_cell(cell, t, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""; pp = cell.paragraphs[0]; pp.alignment = align; pp.paragraph_format.line_spacing = 1.15; runs(pp, t, size, bold)
def grow(tbl, n_rows_needed):
    while len(tbl.rows) < n_rows_needed + 1: tbl._tbl.append(copy.deepcopy(tbl.rows[-1]._tr))
def pageref(cell, bm, size=14):
    cell.text = ""; pp = cell.paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER; cached = str(BM.get(bm, ""))
    def fld(t, instr=None):
        r = pp.add_run(); r.font.size = Pt(size)
        if t in ("begin", "separate", "end"): fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), t); r._r.append(fc)
        else: it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r._r.append(it)
    fld("begin"); fld("instr", f" PAGEREF {bm} \\h "); fld("separate"); rr = pp.add_run(cached); rr.font.size = Pt(size); fld("end")

OV = ROOT / "mining_cache" / "report_pages.json"
if OV.exists(): BM.update({k: int(v) for k, v in json.load(open(OV)).items() if v})
json.dump(BM_TEXT, open(ROOT / "mining_cache" / "report_bm_text.json", "w"))
# lists of figures / tables / abbreviations
T = doc.tables
grow(T[1], len(FIGS)); grow(T[3], len(TABS))
for i, (n, cap, bm) in enumerate(FIGS): row = T[1].rows[i + 1]; fill_cell(row.cells[0], str(n)); fill_cell(row.cells[1], cap, align=WD_ALIGN_PARAGRAPH.LEFT); pageref(row.cells[2], bm, 12)
for i, (n, cap, bm) in enumerate(TABS): row = T[3].rows[i + 1]; fill_cell(row.cells[0], str(n)); fill_cell(row.cells[1], cap, align=WD_ALIGN_PARAGRAPH.LEFT); pageref(row.cells[2], bm, 12)
ABBR = [("API", "Application Programming Interface"), ("AUC", "Area Under the Curve"), ("CART", "Classification and Regression Tree"), ("CI", "Confidence Interval"), ("CSV", "Comma-Separated Values"), ("DOT", "(U.S.) Department of Transportation"), ("DW", "Data Warehouse"), ("ETL", "Extract, Transform, Load"),
        ("GHCN-D", "Global Historical Climatology Network - Daily"), ("IATA", "International Air Transport Association"), ("KNN", "K-Nearest Neighbours"), ("MLP", "Multilayer Perceptron"), ("NAS", "National Air System"), ("NOAA", "National Oceanic and Atmospheric Administration"),
        ("OLAP", "Online Analytical Processing"), ("PR-AUC", "Area Under the Precision-Recall Curve (average precision)"), ("PSO", "Particle Swarm Optimisation"), ("ROC-AUC", "Area Under the Receiver Operating Characteristic Curve"), ("SHAP", "SHapley Additive exPlanations"), ("T-60", "Sixty minutes before scheduled departure")]
grow(T[2], len(ABBR))
for i, (a, b) in enumerate(ABBR): row = T[2].rows[i + 1]; fill_cell(row.cells[0], a); fill_cell(row.cells[1], b, align=WD_ALIGN_PARAGRAPH.LEFT)

# table of contents (the template's own table): page numbers become PAGEREF fields
TOC = {"INTRODUCTION": "ch1", "1.1": "s11", "1.2": "s12", "1.3": "s13", "1.4": "s14", "SYSTEM ANALYSIS": "ch2", "2.1": "s21", "2.2 ": "s22", "2.2.1": "s221", "REQUIREMENT SPECIFICATION": "ch3", "3.1": "s31", "3.2": "s32", "SYSTEM DESIGN SPECIFICATION": "ch4", "4.1": "s41", "4.2": "s42", "4.3": "s43",
       "RESULT AND CONCLUSION": "ch5", "APPENDICES": "ch6", "6.1": "s61", "6.2": "s62", "REFERENCES": "ch7", "7.1": "s71", "7.2": "s72", "7.3": "s73"}
for row in T[0].rows[1:]:
    t1 = re.sub(r"\s+", " ", row.cells[1].text).strip(); t0 = row.cells[0].text.strip(); key = None
    for k, v in TOC.items():
        if t1.startswith(k.strip()) or (t0 and t1.upper().startswith(k.strip().upper()) and not re.match(r"\d", k)): key = v; break
    if t1.startswith("ABSTRACT"): key = "abstract"
    if t1.startswith("2.2.1"): key = "s221"
    elif t1.startswith("2.2"): key = "s22"
    if key and key in BM:
        lead = 0
        for q in row.cells[1].paragraphs:
            if q.text.strip() == "": lead += 1
            else: break
        pageref(row.cells[2], key, 14)
        for _ in range(lead): ne = copy.deepcopy(row.cells[2].paragraphs[0]._p); [c.getparent().remove(c) for c in list(ne) if c.tag == qn("w:r")]; row.cells[2].paragraphs[0]._p.addprevious(ne)
# the template row holding "ABSTRACT / LIST OF FIGURES" and the rows for the two other lists
for row in T[0].rows[1:]:
    t1 = re.sub(r"\s+", " ", row.cells[1].text).strip()
    if t1.startswith("ABSTRACT"): fill_cell(row.cells[2], "", 14); cp = row.cells[2].paragraphs[0]; cp.add_run(str(BM["abstract"])).font.size = Pt(14); cp.add_run().add_break(); cp.add_run(str(BM["lof"])).font.size = Pt(14)
    if t1.startswith("LIST OF ABBREVIATIONS"): pageref(row.cells[2], "loa", 14)
    if t1.startswith("LIST OF TABLES") or t1.startswith("LIST OF  TABLES"): pageref(row.cells[2], "lot", 14)
# chapter 5 sub-entries and appendix wording follow the template; rows for 5.x are left as in the template

# ask Word to refresh page-number fields when the file is opened
st = doc.settings.element; uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true")
for tag in ("w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars", "w:rsids"):
    node = st.find(qn(tag))
    if node is not None: node.addprevious(uf); break
else: st.append(uf)
out = ROOT / "PROJECT_REPORT.docx"; doc.save(out); print("saved", out, "| figures", len(FIGS), "tables", len(TABS), "| est. last page", est["page"], "| bookmarks", len(BM))
