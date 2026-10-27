"""Builds FINAL_REPORT.docx (Word) from the verified result files/tables of Steps 0-5 and the base-paper comparison."""
from pathlib import Path
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt, numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent; R = ROOT / "results"
# ---- comparison chart vs the base paper ----------------------------------------------------------------------
sc = ["Day-ahead", "T-60\n(previous-flight delay)", "Post-departure\n(departure delay)"]
paper_acc, ours_acc = [69.5, 70.9, 89.9], [62.9, 79.2, 88.4]; paper_f2, ours_f2 = [57.2, 58.1, 80.3], [52.5, 61.0, 80.0]
ours_acc_f1 = [69.8, 87.0, 93.1]
fig, axs = plt.subplots(1, 2, figsize=(11, 4.2)); x = np.arange(3); w = 0.27
axs[0].bar(x - w, paper_acc, w, label="Base paper (XGBoost)", color="#999999"); axs[0].bar(x, ours_acc, w, label="Ours, recall matched to paper", color="#2a7ab0")
axs[0].bar(x + w, ours_acc_f1, w, label="Ours, F1-tuned threshold", color="#6aa84f"); axs[0].set_title("Accuracy (%)"); axs[0].set_ylim(50, 100)
axs[1].bar(x - w / 2, paper_f2, w, label="Base paper (XGBoost)", color="#999999"); axs[1].bar(x + w / 2, ours_f2, w, label="Ours, recall matched to paper", color="#2a7ab0"); axs[1].set_title("F2 score (%)"); axs[1].set_ylim(40, 90)
for a in axs: a.set_xticks(x); a.set_xticklabels(sc, fontsize=8); a.spines[["top", "right"]].set_visible(False)
for a, vals in zip([axs[0]] * 3 + [axs[1]] * 2, [(x - w, paper_acc), (x, ours_acc), (x + w, ours_acc_f1), (x - w / 2, paper_f2), (x + w / 2, ours_f2)]):
    for xi, v in zip(*vals): a.text(xi, v + .4, f"{v:.1f}", ha="center", fontsize=7)
axs[0].legend(fontsize=7, loc="upper left"); fig.tight_layout(); CMP = R / "paper_comparison.png"; fig.savefig(CMP, dpi=160); plt.close(fig)

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
for s in doc.sections: s.left_margin = s.right_margin = Inches(0.9); s.top_margin = s.bottom_margin = Inches(0.8)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), color); tcPr.append(sh)
def H(t, l=1): doc.add_heading(t, l)
def P(t, bold=False, italic=False, size=None):
    p = doc.add_paragraph(); r = p.add_run(t); r.bold, r.italic = bold, italic
    if size: r.font.size = Pt(size)
    return p
def B(t):
    p = doc.add_paragraph(style="List Bullet"); 
    if "**" in t:
        parts = t.split("**")
        for i, s in enumerate(parts): p.add_run(s).bold = (i % 2 == 1)
    else: p.add_run(t)
def T(rows, widths=None, hl=None, size=9):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = ""; r = c.paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
            if i == 0: r.bold = True; shade(c, "D9E2F3")
            elif hl and i in hl: shade(c, "E2F0D9"); r.bold = True
    if widths:
        for row in t.rows:
            for j, w_ in enumerate(widths): row.cells[j].width = Inches(w_)
    doc.add_paragraph()
def IMG(rel, cap, w=6.0):
    p = ROOT / rel if not str(rel).startswith("/") else Path(rel)
    if not p.exists(): return
    doc.add_picture(str(p), width=Inches(w)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = doc.add_paragraph(cap); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.runs[0].italic = True; c.runs[0].font.size = Pt(9)

# ============================================== title =========================================================
t = doc.add_heading("Flight Delay & Airport Operations Intelligence", 0); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
P("Data warehouse, weather-impact analysis and delay prediction at T-60 and day-ahead — final project report", italic=True).alignment = WD_ALIGN_PARAGRAPH.CENTER
P("Data: 2015 US DOT on-time performance (5,819,079 flights) + NOAA GHCN-Daily weather.  Base paper: Schösser & Schönberger (2022), Promet – Traffic&Transportation 34(6), 825–838.", size=9).alignment = WD_ALIGN_PARAGRAPH.CENTER

H("1. Problem statement and relevance")
P("Problem. Flight delays and cancellations are driven by several entangled factors: weather, airline operations, airport congestion and the way a late aircraft carries its delay into its next flights. Reports usually blame weather or single causes without separating them. We ask: (1) how much does weather really change the chance that a flight arrives 15 or more minutes late (the official DOT definition of a delay) and the chance it is cancelled; and (2) can we predict, one hour before departure (T-60) and a day ahead, whether a flight will be delayed, and how much do weather and the aircraft's previous leg contribute to that prediction?", bold=False)
P("Relevance.")
B("Passengers and airlines lose time and money to delays; accurate early warnings let airlines re-time crews and aircraft, and let airports prepare gates and ground staff.")
B("Separating weather from other causes shows where operational changes (turnaround buffers, schedule padding, carrier practices) can help and where only forecasting helps. In our data DOT's weather code explains only a fraction of the extra delay on severe-weather days.")
B("Predicting at two moments mirrors real decisions: day-ahead planning (schedule and forecast only) and an operational T-60 alert (adds the aircraft's previous leg).")
B("The study extends a published base paper (Schösser & Schönberger, 2022) with a stricter time-based evaluation, all airports, more algorithms and a dashboard.")
P("Objectives. O1 build a dimensional data warehouse and ETL; O2 OLAP analysis; O3 root-cause and weather-impact quantification; O4 delay propagation along an aircraft's day; O5 predictive mining at T-60 and day-ahead with explainability; O6 airport clustering and cancellation association rules; plus an interactive dashboard and a live demo.")
H("2. Working prototype and implementation")
P("The prototype is a complete, reproducible pipeline from raw files to a dashboard. Every stage runs from a script and reads or writes the DuckDB warehouse.")
T([["Layer", "What was built", "Where"],
   ["Data warehouse", "Star schema (fact_flight + date, airline, airport, route and weather dimensions), 5,819,079 flights; ETL and verification", "etl_pipeline.py, verify_dw.py, mining/build_duckdb.py, flight_dw.duckdb"],
   ["OLAP", "Roll-up, drill-down, slice, dice and pivot by time, airline and airport", "mining/01_olap_export.py, results/olap/"],
   ["Analysis", "Weather impact (controlled and adjusted), cause attribution, delay propagation, data audit", "mining/00_audit.py, 02_root_cause_weather.py, 03_delay_propagation.py"],
   ["Machine learning", "11 models + neural network + ensemble at T-60 and day-ahead; PSO and grid tuning; SHAP explainability; leakage assertion", "mining/04_*.py, 07_neural_net.py, 08_post_departure.py, 09_pso_tuning_final.py"],
   ["Pattern discovery", "K-Means airport clusters; exhaustive and FP-Growth cancellation association rules", "mining/05_clustering.py, 06_association_rules.py"],
   ["Live demo", "Shows the warehouse, trains nine models live on 200,000 flights and scores real flights with SHAP drivers (about 40 s, offline)", "mining/demo_live.py"],
   ["Dashboard", "Interactive web dashboard: time patterns, weather and causes, delay propagation, model comparison with a T-60/day-ahead toggle, base-paper comparison, clusters and rules", "Local multi-page portal (bash run_dashboard.sh, http://localhost:8765): separate pages for T-60, Day-ahead and Post-departure (leaderboard, precision-recall curves, alert slider, weather and drivers, base-paper comparison), a live flight predictor using the saved model with SHAP reasons, an overall metric-evaluation page (scenario comparison, model confidence intervals and heatmap, precision-recall, ROC and lift curves, threshold sweep, base-paper difference charts, tuning), data exploration, light/dark themes and CSV export"]], [1.1, 3.6, 2.1], size=8)
P("How to run: install the pinned packages (requirements.txt), run the scripts in README.md in order (each reads flight_dw.duckdb and writes to results/), and start the demo with 'python mining/demo_live.py --stage all'. Safety: the warehouse is opened read-only by every analysis script, a verified backup with SHA-256 checksum is kept, and a leakage assertion fails the run if any feature uses information from after T-60.")
P("Quality controls: a time-based train/validation/test split, class-weighted training with the threshold tuned on validation data only, multiple seeds, bootstrap confidence intervals, matched-training-size comparisons and paired significance tests.", size=10)
H("3. Summary")
B("**Weather matters.** Severe origin weather raises the share of flights arriving 15+ min late from 17.0% to 40.4% (+23.4 percentage points after controlling for airport and month; odds ratio 3.6) and multiplies cancellations about 8x. Adding weather features lifts prediction PR-AUC by +0.04 at T-60 and by +0.10 day-ahead (95% CIs exclude zero).")
B("**Delays propagate.** 60.8% of 15+ min departure delays follow a late inbound aircraft; knowing the previous leg lifts PR-AUC from 0.36 (day-ahead) to 0.64 (T-60).")
B("**Best model:** an ensemble of LightGBM, XGBoost and a neural network: PR-AUC 0.645 / ROC-AUC 0.824 at T-60 (3.6x lift over the 17.95% base rate) and 0.360 / 0.703 day-ahead, on a strict time-based test (Nov–Dec).")
B("**Versus the base paper:** we are higher at T-60 (accuracy 79.2% vs 70.9%, F2 61.0% vs 58.1% at the same recall) and on post-departure accuracy at our own threshold (93.1% vs 89.9%); roughly tied on post-departure F2 (80.0% vs 80.3%); behind on the day-ahead task. Our evaluation is stricter (time-based split, all airports) and reports precision, PR-AUC and ROC-AUC with confidence intervals, which the paper does not.")

H("4. Data and warehouse")
P("Delayed = arrival delay >= 15 minutes (DOT rule; the same definition as the base paper). Two prediction moments are studied: day-ahead (schedule, calendar, airline/airport/route and observed same-day weather as a forecast proxy) and T-60 (the same plus the aircraft's previous-leg status and delay, known 60 minutes before scheduled departure). A leakage assertion checks that no feature uses information after T-60.")
T([["Item", "Value"], ["Flights", "5,819,079 (2015); October excluded from airport-level work (airport codes unmapped)"], ["Warehouse", "DuckDB star schema: fact_flight + dim_date, dim_airline, dim_airport, dim_route, dim_weather"],
   ["Split", "Train Jan–Aug (3,836,893) / validation Sep (462,153) / test Nov–Dec (932,084)"], ["Test base rate", "17.95% delayed"], ["Overall on-time", "81.39% (arrival delay < 15 min)"]], [1.5, 5.0])

H("5. Descriptive results")
H("5.1 Time patterns (OLAP)", 2)
B("Delay builds through the day: on-time falls from 92.7% (05:00 departures) to 73.7% (20:00); evening departures are 2.4x as likely to leave 15+ min late as morning ones (27.0% vs 11.3%).")
B("Worst months: February (76.65% on-time, 4.78% cancelled) and June (76.52%); best: September (87.00%). Saturday is the best weekday (84.01% on-time); Monday has the highest cancellation rate (2.43%).")
B("Airlines: Hawaiian 88.67% on-time, Alaska 86.96%, Delta 86.44% … Frontier 73.84%, Spirit 70.29%. American Eagle cancels 5.1% of flights, 3x the 1.54% network rate.")
IMG("results/olap/by_dep_hour.png", "Figure 1. On-time performance by scheduled departure hour.", 5.6)
IMG("results/olap/by_month.png", "Figure 2. Monthly delay and cancellation profile.", 5.6)

H("5.2 Root causes and weather impact", 2)
T([["Delay cause (flights >= 15 min late)", "Delay minutes", "Share"], ["Late aircraft", "24,961,931", "39.8%"], ["Airline", "20,172,956", "32.2%"], ["Air system (NAS)", "14,335,762", "22.9%"], ["Weather", "3,100,233", "4.9%"], ["Security", "80,985", "0.1%"]], [3, 1.6, 1.0])
T([["Outcome (adjusted for airport, month, hour, airline)", "Moderate weather OR (95% CI)", "Severe weather OR (95% CI)"], ["Delayed (arrival >= 15 min)", "1.98 (1.90–2.05)", "3.60 (3.35–3.87)"], ["Cancelled", "2.55 (2.26–2.88)", "7.70 (6.57–9.03)"]], [3, 1.8, 1.8])
B("Severe origin weather: delay rate 40.4% vs 17.0% on clear days (+23.4 pp controlled for airport and month), +26 min average arrival delay, cancellation rate 9.0% vs 1.1%. Counting cancellations as disruptions, the gap rises to +27.5 pp.")
B("Weather causes 54.8% of cancellations (47,691 of 87,076). DOT's 'weather' code captures only about a fifth (20%) of the excess delay on severe days; air-system delay rises by about as much.")
IMG("results/root_cause_weather/2a_cause_share_overall.png", "Figure 3. Share of delay minutes by DOT cause.", 5.2)
IMG("results/root_cause_weather/2c_raw_vs_controlled.png", "Figure 4. Weather effect: raw vs controlled for airport and month.", 5.6)
IMG("results/root_cause_weather/2b_cancel_reasons.png", "Figure 5. Cancellation reasons.", 5.6)

H("5.3 Delay propagation", 2)
T([["Previous-leg arrival delay", "Mean departure delay (min)", "Departing >= 15 min late"], ["< -15 (early)", "0.0", "6.6%"], ["0..14", "6.9", "16.5%"], ["15..29", "19.1", "54.1%"], ["30..59", "37.0", "83.2%"], ["60..89", "62.6", "89.1%"]], [2.4, 2.0, 2.0])
B("Mean departure delay rises from 5.0 min on an aircraft's first leg to 16.4 min on its seventh. Planned turns of 90+ minutes absorb about three quarters of delays under an hour; turns under 45 minutes absorb nothing.")
B("Best absorbers: SkyWest (-8.2 min), Hawaiian (-8.0); worst amplifiers: United (+8.8), JetBlue (+5.8), Spirit (+3.1).")
IMG("results/propagation/3A_propagation_curve.png", "Figure 6. Departure delay vs the aircraft's incoming delay.", 5.6)
IMG("results/propagation/3B_buffer_absorption.png", "Figure 7. Delay absorbed as a function of planned turnaround.", 5.6)

H("6. Prediction of delay >= 15 minutes")
H("6.1 Protocol", 2)
P("Eleven model variants were compared on identical splits with class-weighted training and a decision threshold tuned on September (F1). PR-AUC is the primary metric because 82% of flights are on time (always predicting 'on time' scores 82% accuracy). Tuning: grid search on validation data, particle swarm optimisation (as in the base paper) for XGBoost, LightGBM and the neural network. Bootstrap confidence intervals and multiple seeds quantify uncertainty.")
H("6.2 T-60 (operational model) — test set", 2)
T([["Model", "PR-AUC", "ROC-AUC", "Accuracy", "Precision", "Recall", "F1"],
   ["Ensemble (LightGBM + XGBoost + MLP)", "0.6446", "0.8242", "86.96%", "72.4%", "44.3%", "0.550"], ["XGBoost (native categoricals, tuned)", "0.6409", "0.8214", "85.21%", "60.0%", "52.7%", "0.562"],
   ["LightGBM (PSO-tuned)", "0.6424", "0.8220", "85.09%", "59.4%", "53.5%", "0.563"], ["LightGBM", "0.6388", "0.8202", "84.82%", "58.4%", "53.8%", "0.560"],
   ["Neural network (MLP)", "0.6338", "0.8173", "85.89%", "64.1%", "48.7%", "0.554"], ["Random Forest (1M rows)", "0.6189", "0.8128", "84.92%", "59.3%", "51.2%", "0.550"],
   ["Decision Tree", "0.6041", "0.7929", "85.75%", "64.8%", "45.3%", "0.533"], ["Logistic Regression", "0.5502", "0.7806", "82.76%", "52.1%", "49.4%", "0.507"],
   ["KNN (200k rows)", "0.4784", "0.7400", "–", "51.2%", "41.7%", "0.460"], ["Naive Bayes (categorical)", "0.4593", "0.7318", "–", "51.2%", "36.9%", "0.429"], ["Naive Bayes (Gaussian)", "0.3624", "0.7198", "–", "34.0%", "60.7%", "0.436"]], [2.5, .7, .7, .75, .75, .65, .5], hl=[1], size=8)
IMG("results/model_comparison/model_comparison_t60.png", "Figure 8. T-60 model comparison (PR-AUC).", 5.6)
IMG("results/model_comparison/pr_curves_t60.png", "Figure 9. Precision–recall curves at T-60.", 5.6)
H("6.3 Day-ahead — test set", 2)
T([["Model", "PR-AUC", "ROC-AUC", "Accuracy", "Precision", "Recall", "F1"],
   ["Ensemble (LightGBM + XGBoost + MLP)", "0.3597", "0.7028", "75.02%", "34.7%", "44.5%", "0.390"], ["XGBoost (native categoricals, tuned)", "0.3559", "0.6996", "69.81%", "30.9%", "55.3%", "0.397"],
   ["LightGBM", "0.3522", "0.6986", "70.31%", "31.2%", "54.1%", "0.396"], ["Random Forest", "0.3416", "0.6934", "68.25%", "29.8%", "56.7%", "0.391"], ["Neural network (MLP)", "0.3373", "0.6887", "73.38%", "32.5%", "44.8%", "0.377"],
   ["Logistic Regression", "0.3196", "0.6704", "69.85%", "29.7%", "49.6%", "0.371"], ["Naive Bayes (Gaussian)", "0.2877", "0.6529", "49.48%", "22.8%", "75.9%", "0.350"], ["Decision Tree", "0.2849", "0.6366", "73.07%", "29.7%", "36.6%", "0.328"],
   ["Majority class (dummy)", "0.1795", "0.5000", "82.05%", "0%", "0%", "0.000"]], [2.5, .7, .7, .75, .75, .65, .5], hl=[1], size=8)
IMG("results/model_comparison/model_comparison_dayahead.png", "Figure 10. Day-ahead model comparison (PR-AUC).", 5.6)

H("6.4 Weather impact on prediction (tuned models, weather features ON vs OFF)", 2)
T([["Task", "Model", "PR-AUC without weather", "PR-AUC with weather", "Gain (95% CI)"], ["T-60", "XGBoost (PSO)", "0.599", "0.641", "+0.042 [0.041, 0.043]"], ["T-60", "LightGBM (PSO)", "0.599", "0.642", "+0.043 [0.043, 0.044]"],
   ["Day-ahead", "XGBoost (PSO)", "0.247", "0.344", "+0.097 [0.095, 0.099]"], ["Day-ahead", "LightGBM (PSO)", "0.244", "0.347", "+0.103 [0.101, 0.105]"]], [1.0, 1.5, 1.5, 1.4, 1.6])
P("Weather matters most when nothing is known about the aircraft's earlier legs (day-ahead), where it roughly doubles the model's lift over the base rate; at T-60 the previous leg already carries most of the information. Weather accounts for 11.7% of LightGBM's gain, behind predecessor status/delay (42%) and airline/airport/route (24%). The weather used is observed same-day weather, an upper bound on what a real forecast would provide.")
IMG("results/model_comparison/ablation_lightgbm.png", "Figure 11. LightGBM ablation: effect of weather, propagation features and history.", 5.6)
IMG("results/explain/shap_summary.png", "Figure 12. SHAP summary: most influential features at T-60.", 5.6)
H("6.5 Hyperparameters found by particle swarm optimisation", 2)
T([["Task", "Model", "Tuned values", "Test PR-AUC"],
   ["T-60", "XGBoost", "max_depth 9, lr 0.055, min_child_weight 40, colsample 0.48, gamma 0.11, subsample 0.99", "0.6407"],
   ["T-60", "LightGBM", "num_leaves 164, lr 0.05, min_child_samples 674, feature_fraction 0.40, lambda_l2 0.01, bagging 1.0", "0.6424"],
   ["Day-ahead", "XGBoost", "max_depth 9, lr 0.05, min_child_weight 56, colsample 0.40, gamma 1.02, subsample 1.0", "0.3437"],
   ["Day-ahead", "LightGBM", "num_leaves 78, lr 0.098, min_child_samples 2000, feature_fraction 0.77, lambda_l2 0.81, bagging 1.0", "0.3469"],
   ["T-60", "MLP", "width 676, dropout 0.10, lr 0.0076, weight decay 8.8e-5, batch 1024", "0.6254"], ["Day-ahead", "MLP", "width 781, dropout 0.20, lr 0.0038, weight decay 2.1e-4, batch 1024", "0.3327"]], [0.8, 0.8, 4.0, 0.9], size=8)
P("Tuning gave at most +0.004 PR-AUC (T-60 LightGBM) and hurt day-ahead LightGBM and both MLPs, because settings chosen on a 600k-row sample did not transfer to full-data training. The earlier configurations remain the reference results above.", italic=True, size=9)

H("7. Pattern discovery")
H("7.1 Airport clusters", 2)
T([["Cluster", "Airports", "Delayed %", "Cancel %", "Weather gap (pp)", "Examples"], ["Mid-size, punctual", "35", "15.9", "0.94", "7.6", "PHX, SEA, SLC, FLL, SAN, TPA"], ["Large hubs, congested", "24", "21.0", "1.77", "16.9", "ATL, ORD, DFW, DEN, LAX, SFO"], ["Smaller, weather-exposed", "54", "18.1", "2.33", "10.7", "BNA, CLE, RDU, MKE, IND, CMH"]], [1.6, .7, .8, .8, 1.0, 2.0])
P("K-Means on 113 airports (k = 3): silhouette 0.24 (soft boundaries); agrees with Ward hierarchical clustering (ARI 0.82); stable under resampling (mean ARI 0.88). The weather gap is the rise in delay rate in adverse vs clear weather.", size=9)
IMG("results/clustering/airport_clusters_pca.png", "Figure 13. Airport clusters (PCA projection).", 5.2)
IMG("results/clustering/cluster_profile_heatmap.png", "Figure 14. Cluster profiles (z-scores).", 5.4)
H("7.2 Cancellation association rules", 2)
T([["Rule (antecedent -> CANCELLED)", "Flights", "Cancelled", "Confidence", "Lift"], ["airline = MQ & origin snow & destination severe weather", "2,795", "1,274", "45.6%", "27.7"], ["origin severe weather & origin high wind", "3,542", "1,568", "44.3%", "26.9"],
   ["origin snow & origin high wind", "3,359", "1,469", "43.7%", "26.6"], ["origin severe weather & freezing & destination severe weather", "5,475", "1,848", "33.8%", "20.5"], ["origin high wind & origin freezing", "5,267", "1,395", "26.5%", "16.1"]], [3.4, .7, .8, .9, .6])
P("Exhaustive mining of 5.30M flights (87,076 cancelled, base rate 1.64%); lift = P(cancelled | antecedent) / 1.64%. Among cancelled flights (FP-Growth), clear weather at both airports points to an airline-caused cancellation (confidence 55%, lift 1.98) and origin snow to a weather-caused one (54%). These are associations, not causes.", size=9)
IMG("results/association_rules/top_cancellation_rules.png", "Figure 15. Top non-redundant cancellation rules.", 5.6)

H("8. Comparison with the base paper")
P("Schösser & Schönberger (2022) predict arrival delay >= 15 min on the same 2015 DOT data with Random Forest, XGBoost and a neural network, in three scenarios: static features and weather (set 1), plus the previous flight's delay (set 1+2) and plus departure delay (set 1+3). They report accuracy, recall and F2 (beta = 2). We map their sets to our day-ahead, T-60 and a post-departure scenario that we added for this comparison (departure delay is only known after take-off, so it is not an operational model).")
P("Because the paper trains on undersampled data (which raises recall and lowers accuracy), we show our models at two operating points: our F1-tuned threshold, and a lowered threshold giving recall near the paper's (set on validation data; test recall is only approximately matched).", size=9.5)
H("8.1 Accuracy, precision, recall and F2 side by side (%)", 2)
T([["Scenario", "System", "Accuracy", "Precision", "Recall", "F2"],
   ["Day-ahead", "Paper XGBoost", "69.5", "not reported", "68.2", "57.2"], ["", "Paper Random Forest", "68.7", "n/r", "67.2", "56.2"], ["", "Paper Neural Network", "66.1", "n/r", "65.8", "54.0"],
   ["", "Ours XGBoost, F1 threshold", "69.8", "30.9", "55.3", "47.8"], ["", "Ours Ensemble, recall ~67%", "62.9", "27.9", "67.3", "52.5"],
   ["T-60", "Paper XGBoost", "70.9", "n/r", "68.5", "58.1"], ["", "Paper Random Forest", "70.1", "n/r", "67.2", "56.9"], ["", "Paper Neural Network", "66.6", "n/r", "66.5", "54.7"],
   ["", "Ours XGBoost, F1 threshold", "85.2", "60.0", "52.7", "54.1"], ["", "Ours Ensemble, F1 threshold", "87.0", "72.4", "44.3", "48.0"], ["", "Ours Ensemble, recall ~67%", "79.2", "44.7", "67.1", "61.0"], ["", "Ours XGBoost, recall ~75%", "72.4", "36.9", "75.3", "62.3"],
   ["Post-departure", "Paper XGBoost", "89.9", "n/r", "83.4", "80.3"], ["", "Paper Random Forest", "90.2", "n/r", "81.7", "79.4"], ["", "Paper Neural Network", "86.5", "n/r", "85.3", "78.8"],
   ["", "Ours XGBoost, F1 threshold", "93.1", "84.4", "75.8", "77.3"], ["", "Ours LightGBM, recall ~86%", "88.4", "62.9", "85.8", "80.0"]], [1.1, 2.1, .8, .9, .7, .6], hl=[11, 12, 16], size=8)
H("Post-departure scenario in detail (paper feature set 1+3)", 3)
T([["", "Accuracy", "Recall", "F2", "PR-AUC", "ROC-AUC"],
   ["Paper XGBoost, tuned", "89.9%", "83.4%", "80.3%", "not reported", "not reported"], ["Paper Random Forest, tuned", "90.2%", "81.7%", "79.4%", "not reported", "not reported"],
   ["Paper Neural Network, tuned", "86.5%", "85.3%", "78.8%", "not reported", "not reported"],
   ["Ours XGBoost, F1-tuned threshold", "93.1%", "75.8%", "77.3%", "0.880", "0.944"], ["Ours LightGBM, F1-tuned threshold", "93.1%", "75.9%", "77.3%", "0.879", "0.943"],
   ["Ours XGBoost, threshold set for ~86% recall", "87.9%", "86.3%", "79.9%", "0.880", "0.944"], ["Ours LightGBM, threshold set for ~86% recall", "88.4%", "85.8%", "80.0%", "0.879", "0.943"]], [2.8, .8, .7, .7, .9, .9], hl=[4, 5], size=8.5)
P("Precision (not reported by the paper): ours 84.4% (XGBoost) and 83.9% (LightGBM) at the F1 threshold; 61.7% and 62.9% at ~86% recall. PR-AUC and ROC-AUC do not depend on the threshold. The paper does not report them, so there is nothing to compare them against; they show how well our models rank flights.", size=9)
IMG(CMP, "Figure 16. Accuracy and F2 against the base paper, three scenarios (XGBoost in the paper; our best model at matched recall and at our own F1 threshold).", 6.2)
H("8.2 Where our results are higher", 2)
T([["Comparison", "Paper", "Ours", "Difference"],
   ["T-60 accuracy at the same recall (~67–68%)", "70.9%", "79.2% (ensemble)", "+8.3 points"], ["T-60 F2 at the same recall", "58.1%", "61.0% (ensemble)", "+2.9 points"],
   ["T-60 accuracy at 75% recall (higher than paper)", "70.9%", "72.4% (XGBoost)", "+1.5 points"], ["T-60 F2 at 75% recall", "58.1%", "62.3% (XGBoost)", "+4.2 points"],
   ["Post-departure accuracy (our F1 threshold)", "89.9%", "93.1% (XGBoost)", "+3.2 points"], ["Day-ahead accuracy (our F1 threshold)", "69.5%", "69.8% (XGBoost)", "+0.3 points"],
   ["Effect of the previous-flight delay (PR-AUC, day-ahead to T-60)", "accuracy +1.4 points", "PR-AUC 0.36 to 0.64", "much larger gain"],
   ["Algorithms compared", "3", "11 variants + ensemble", "wider coverage"], ["Metrics reported", "accuracy, recall, F2", "+ precision, F1, PR-AUC, ROC-AUC, confidence intervals", "richer"],
   ["Evaluation design", "random 75:25 split, 30 busiest airports (2.23M flights)", "time-based Nov–Dec test, all airports (5.82M flights)", "stricter, no date leakage"],
   ["Scope", "prediction only", "warehouse, OLAP, root cause, propagation, clustering, association rules", "broader"]], [2.6, 1.4, 1.9, 1.0], hl=[1, 2, 3, 4, 5], size=8)
H("8.3 Where the paper is ahead or level (stated plainly)", 2)
B("Day-ahead task: the paper's F2 is higher (57.2% vs our 52.5% at matched recall) and its accuracy at matched recall is higher (69.5% vs 62.9%). Its random split and 30-airport subset make that task easier than our time-based test.")
B("Post-departure: F2 is level (80.0% vs 80.3%); at matched recall (~86% vs 83.4%) our accuracy is about 2 points lower (88.4% vs 89.9%). The accuracy advantage above holds only at our own, lower-recall threshold.")
B("The paper tunes with grid search plus particle swarm optimisation, as we now do; our swarm tuning added at most +0.004 PR-AUC.")
B("Our 'matched recall' rows use thresholds set on September validation data, so test recall lands within a few points of the target; the comparison is approximate.")
H("8.4 Agreement with the paper's findings", 2)
B("Tree-based ensembles (Random Forest, XGBoost, and here LightGBM) beat the neural network, which was slightly behind in both studies.")
B("The previous flight's delay helps, and departure delay helps most; the paper's conclusion that the best feature set depends on the desired prediction time is confirmed.")
B("Accuracy alone is misleading with an ~80% on-time majority: both studies emphasise recall and F-scores.")
B("SHAP shows scheduled departure time, weather, airline and the previous-flight delay among the most influential features, as in the paper.")

H("9. Conclusions")
B("Weather measurably drives both delay and cancellation: +23 pp delay rate and about 8x cancellation on severe-weather days, and +0.04 (T-60) to +0.10 (day-ahead) PR-AUC when weather features are added.")
B("A T-60 model that knows the aircraft's previous leg predicts delays at 3.6x the base rate; the ensemble reaches PR-AUC 0.645 and exceeds the base paper's accuracy and F2 at the same recall.")
B("The strongest signals are operational (previous leg, time of day, airline/airport), weather is next, and planned turnaround buffers absorb delay propagation.")
H("Limitations", 2)
B("Weather is same-day observed weather, a proxy for a forecast. One test period (Nov–Dec). Random Forest and KNN were trained on samples; the ensemble's small lead over XGBoost (+0.004 PR-AUC) has not been tested for significance. Clusters are soft (silhouette 0.24) and the rules are associations. The base-paper comparison is approximate because of different splits, data subsets and class-balancing.")
doc.save(ROOT / "FINAL_REPORT.docx"); print("saved", ROOT / "FINAL_REPORT.docx")
